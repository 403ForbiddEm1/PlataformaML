"""Extracción de texto de hojas de vida y vacantes en formato PDF o DOCX."""
import io

from docx import Document
from pypdf import PdfReader

from app.config import MIN_TEXT_CHARS


class ExtractionError(Exception):
    """Error al extraer texto; lleva el código HTTP que corresponde."""

    def __init__(self, message: str, status_code: int = 422):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


def _from_pdf(content: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(content))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as exc:  # archivo dañado o protegido
        raise ExtractionError("No se pudo leer el archivo PDF", 422) from exc


def _from_docx(content: bytes) -> str:
    try:
        doc = Document(io.BytesIO(content))
    except Exception as exc:
        raise ExtractionError("No se pudo leer el archivo DOCX", 422) from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def extract_text(filename: str, content: bytes) -> str:
    """Devuelve el texto del documento. Rechaza formatos no soportados y documentos sin texto procesable."""
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        text = _from_pdf(content)
    elif name.endswith(".docx"):
        text = _from_docx(content)
    else:
        raise ExtractionError("Formato no soportado: solo se aceptan archivos PDF o DOCX", 415)
    text = " ".join(text.split())
    if len(text) < MIN_TEXT_CHARS:
        raise ExtractionError(
            "El documento no contiene texto procesable (por ejemplo, es un documento escaneado)", 422
        )
    return text
